
import streamlit as st
import pandas as pd
from docx import Document
import io

st.set_page_config(page_title="ระบบรายงานผลการเรียน", layout="centered")

st.title("📑 ระบบสร้างรายงานผลการเรียน")

# อัปโหลดไฟล์
excel_file = st.file_uploader("📥 เลือกไฟล์ Excel (เกรดนักเรียน)", type=["xlsx"])
template_file = st.file_uploader("📄 เลือกไฟล์ Word Template", type=["docx"])

# ตัวเลือกทั่วไป
report_type = st.radio("เลือกรูปแบบการสร้างรายงาน", ["ทั้งหมด", "เฉพาะนักเรียนที่เลือก"])
output_format = st.selectbox("เลือกรูปแบบไฟล์", ["docx"])  # รองรับเฉพาะ docx ในเวอร์ชันนี้

# รายชื่อที่เลือก (หากเป็นรายคน)
student_ids = st.text_input("กรอกเลขประจำตัวนักเรียน (คั่นด้วย comma)")

# รายละเอียดหัวเอกสาร
st.subheader("📝 รายละเอียดหัวรายงาน")
term = st.text_input("ภาคเรียนที่", value="2")
year = st.text_input("ปีการศึกษา", value="2566")
grade = st.text_input("ระดับชั้น", value="มัธยมศึกษาปีที่ 1/9")
program = st.text_input("โปรแกรม", value="SME แสงทอง")

def replace_placeholders(doc, replacements):
    for p in doc.paragraphs:
        for key, val in replacements.items():
            if key in p.text:
                inline = p.runs
                for i in range(len(inline)):
                    if key in inline[i].text:
                        inline[i].text = inline[i].text.replace(key, val)
    return doc

if st.button("🚀 สร้างรายงาน"):
    if excel_file and template_file:
        df = pd.read_excel(excel_file)
        df['เลขประจำตัวนักเรียน'] = df['เลขประจำตัวนักเรียน'].astype(str)

        if report_type == "ทั้งหมด":
            selected_students = df
        else:
            ids = [x.strip() for x in student_ids.split(",") if x.strip()]
            selected_students = df[df['เลขประจำตัวนักเรียน'].isin(ids)]

        for _, student in selected_students.iterrows():
            doc = Document(template_file)

            # แทนค่าหัวกระดาษ
            header_replacements = {
                "ภาคเรียนที่ 2": f"ภาคเรียนที่ {term}",
                "ปีการศึกษา 2566": f"ปีการศึกษา {year}",
                "มัธยมศึกษาปีที่ 1/9": grade,
                "SME แสงทอง": program
            }

            doc = replace_placeholders(doc, header_replacements)

            # สร้าง dict สำหรับแทนที่ค่า
            replace_dict = {
                '«title»': str(student.get('คำนำหน้า', '')),
                '«name»': str(student.get('ชื่อ', '')),
                '«last»': str(student.get('นามสกุล', '')),
                '«id»': str(student.get('เลขประจำตัวนักเรียน', '')),
                '«gt1»': str(student.get('ท21102', '')),
                '«gt2»': str(student.get('ค21102', '')),
                '«gt3»': str(student.get('ว21102', '')),
                '«gt4»': str(student.get('ส21103', '')),
                '«gt5»': str(student.get('ส21104', '')),
                '«gt6»': str(student.get('พ21102', '')),
                '«gt7»': str(student.get('ศ21102', '')),
                '«gt8»': str(student.get('ง21102', '')),
                '«gt9»': str(student.get('อ21102', '')),
                '«pt1»': str(student.get('ว21282', '')),
                '«pt2»': str(student.get('ค21202', '')),
                '«pt3»': str(student.get('ว21204', '')),
                '«pt4»': str(student.get('อ21208', '')),
                '«pt5»': str(student.get('อ21210', '')),
                '«pt6»': str(student.get('อ21212', '')),
                '«pt7»': str(student.get('ส21202', '')),
                '«grade2»': str(student.get('GPA', ''))
            }

            doc = replace_placeholders(doc, replace_dict)

            # บันทึกในหน่วยความจำ
            buffer = io.BytesIO()
            doc.save(buffer)
            buffer.seek(0)

            file_name = f"รายงาน_{student['เลขประจำตัวนักเรียน']}_{student['ชื่อ']}.docx"
            st.download_button(
                label=f"📄 ดาวน์โหลด: {file_name}",
                data=buffer,
                file_name=file_name,
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            )
    else:
        st.error("⚠️ กรุณาอัปโหลดไฟล์ Excel และ Template Word ให้ครบถ้วนก่อนสร้างรายงาน")
